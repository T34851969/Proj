"""Document text extraction for .docx, .pdf, .txt, .md files.

Uploads are validated by extension AND magic bytes before parsing, so a
renamed/forge file fails fast with a 400 instead of corrupting the KB.
"""

from __future__ import annotations

from io import BytesIO
from typing import Tuple

from docx import Document as DocxDocument


def parse_docx(file_bytes: bytes) -> str:
    """Extract plain text from a .docx file (paragraphs + table cells)."""
    doc = DocxDocument(BytesIO(file_bytes))
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                paragraphs.append(" | ".join(cells))
    return "\n".join(paragraphs)


def parse_pdf(file_bytes: bytes) -> str:
    """Extract plain text from a .pdf file using pypdf."""
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(file_bytes))
    texts: list[str] = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            texts.append(text.strip())
    return "\n".join(texts)


def parse_txt(file_bytes: bytes) -> str:
    """Decode text file with UTF-8 then GBK fallback."""
    try:
        return file_bytes.decode("utf-8")
    except UnicodeDecodeError:
        return file_bytes.decode("gbk", errors="replace")


def validate_magic_bytes(filename: str, file_bytes: bytes) -> None:
    """Reject files whose content does not match the declared extension."""
    lower_name = filename.lower()
    head = file_bytes[:8]
    if lower_name.endswith(".docx") and not head.startswith(b"PK\x03\x04"):
        raise ValueError("文件内容与 .docx 格式不符(文件可能已损坏或被改名)")
    if lower_name.endswith(".pdf") and not head.startswith(b"%PDF-"):
        raise ValueError("文件内容与 .pdf 格式不符(文件可能已损坏或被改名)")
    if lower_name.endswith((".txt", ".md")) and b"\x00" in file_bytes[:4096]:
        raise ValueError("文件包含二进制内容，无法作为文本解析")


def parse_file(filename: str, file_bytes: bytes) -> Tuple[str, str]:
    """Parse uploaded file and return (filename, text).

    Raises ValueError for unsupported or mismatched file types.
    """
    validate_magic_bytes(filename, file_bytes)
    lower_name = filename.lower()
    if lower_name.endswith(".docx"):
        return filename, parse_docx(file_bytes)
    if lower_name.endswith(".pdf"):
        return filename, parse_pdf(file_bytes)
    if lower_name.endswith(".txt") or lower_name.endswith(".md"):
        return filename, parse_txt(file_bytes)
    raise ValueError(f"不支持的文件类型: {filename}，仅支持 .docx, .pdf, .txt, .md")
